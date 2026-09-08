from __future__ import annotations

import json
from io import BytesIO

from openpyxl import load_workbook

from idx_flow_scanner.providers.block_idx_evidence import download_official_idx_attachment
from idx_flow_scanner.providers.block_idx_financial_history import (
    fetch_financial_report_market,
    parse_financial_report_payload,
)


def clean(value: object) -> object:
    if value is None or isinstance(value, (int, float, bool)):
        return value
    text = str(value).replace("\n", " ").strip()
    return text[:180]


def main() -> None:
    payload = fetch_financial_report_market(2026, "TW2", ticker="BSSR")
    reports = parse_financial_report_payload(payload)
    report = next(r for r in reports if r["ticker"] == "BSSR")
    attachment = next(a for a in report["attachments"] if str(a["file_name"]).lower().endswith(".xlsx"))
    data, digest, content_type = download_official_idx_attachment(str(attachment["file_url"]))
    wb = load_workbook(BytesIO(data), read_only=True, data_only=False)
    telemetry: dict[str, object] = {
        "ticker": "BSSR",
        "report": {k: report[k] for k in ("report_year", "report_period", "report_period_end", "file_modified_at")},
        "file_name": attachment["file_name"],
        "sha256": digest,
        "bytes": len(data),
        "content_type": content_type,
        "sheets": [],
    }
    sheets: list[dict[str, object]] = []
    for ws in wb.worksheets:
        sample: list[list[object]] = []
        nonempty = 0
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 24), min_col=1, max_col=min(ws.max_column, 10), values_only=True):
            values = [clean(v) for v in row]
            if any(v not in (None, "") for v in values):
                sample.append(values)
                nonempty += 1
            if nonempty >= 12:
                break
        sheets.append({
            "title": ws.title,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "sample": sample,
        })
    telemetry["sheets"] = sheets
    print(json.dumps(telemetry, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
